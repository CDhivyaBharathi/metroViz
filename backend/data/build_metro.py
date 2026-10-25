"""One-off: build metro.json from three sources.

    backend/.venv/bin/python backend/data/build_metro.py            # cached osm_raw.json
    backend/.venv/bin/python backend/data/build_metro.py --refresh  # re-download OSM

Sources (PLAN.md D4):
  * OpenStreetMap via Overpass -> stations, line order, coordinates, colours (most current).
    © OpenStreetMap contributors, ODbL 1.0.
  * DMRC static GTFS (Aug 2023) -> official ride times between stations.
  * DMRC station & gate locations spreadsheet -> official station codes, gates/exits.
    Both from Delhi Transport Stack / Delhi Open Transit Data, Transport Department, GNCTD.
    Raw files live in backend/data/raw/ and are NOT committed (OTD terms §7).
Connections GTFS doesn't cover (post-2023 openings) get an estimate fitted on the GTFS times.
"""
import collections
import csv
import itertools
import json
import math
import re
import statistics
import sys
import urllib.parse
import urllib.request
import zlib
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
OSM_RAW = HERE / "osm_raw.json"
GTFS = HERE / "raw" / "dmrc_static_gtfs_v1"
XLSX = HERE / "raw" / "dmrc_station_and_gate_locations.xlsx"
OUT = HERE / "metro.json"

# --- tuning -------------------------------------------------------------------------
MATCH_RADIUS_KM = 0.3   # GTFS/spreadsheet station within this of an OSM station = same station
MIN_MINUTES = 0.5
HOLDOUT_SHARE = 0.3     # share of GTFS-timed connections kept out of the estimate fit, to test it
SEPARATE_SPEED_LINES = {"AIRPORT"}  # genuinely faster; gets its own fitted speed

LINES = {  # OSM name prefix -> (id, display name)
    "Red Line": ("RED", "Red Line"),
    "Yellow Line": ("YELLOW", "Yellow Line"),
    "Blue Line": ("BLUE", "Blue Line"),
    "Green Line": ("GREEN", "Green Line"),
    "Violet Line": ("VIOLET", "Violet Line"),
    "Pink Line": ("PINK", "Pink Line"),
    "Magenta Line": ("MAGENTA", "Magenta Line"),
    "Grey Line": ("GREY", "Grey Line"),
    "Airport Express Line": ("AIRPORT", "Airport Express Line"),
    "Rapid MetroRail Gurgaon": ("RAPID", "Rapid Metro Gurugram"),
}
# Branches run as separate services (you change trains at Yamuna Bank / Ashok Park Main),
# so they get their own line id; otherwise interchange detection would miss the change.
BRANCHES = {  # (line id, substring of OSM route name) -> (branch id, display name)
    ("BLUE", "Vaishali"): ("BLUE_BR", "Blue Line (Vaishali branch)"),
    ("GREEN", "Kirti Nagar"): ("GREEN_BR", "Green Line (Kirti Nagar branch)"),
}
COLOR_FIX = {"gray": "#808080"}  # OSM sometimes uses CSS names

# Same physical interchange, different OSM names. Checked by hand against DMRC's map.
# Deliberately NOT merged (close but not interchanges): ITO/Supreme Court,
# Janpath/Patel Chowk, Nehru Enclave/Nehru Place.
ALIASES = {
    "New Delhi Airport Express Terminal": "New Delhi",
    "Sikanderpur": "Sikandarpur",
}
GTFS_EXCLUDE_PREFIX = "AQUA"  # Noida Metro (NMRC): out of scope (PLAN.md D9)

# Spreadsheet name -> OSM name, for names that differ beyond spelling normalisation and
# whose spreadsheet coordinates are unusable. Checked by hand.
XLSX_ALIASES = {
    "Anand Vihar Isbt": "Anand Vihar",
    "Badarpur Border": "Badarpur",
    "Raja Nahar Singh Balabgarh": "Raja Nahar Singh",
    "Brig. Hoshiar Singh": "Brigadier Hoshiyar Singh",
    "Golf Course": "Noida Golf Course",
    "Guru Teg Bahadur Station": "GTB Nagar",
    "JLN Stadium": "Jawaharlal Nehru Stadium",
    "Mayur Vihar Pocket-1": "Shree Ram Mandir Mayur Vihar",
    "Pitampura": "Madhuban Chowk",
    "Majar Mohit Sharma Rajendra Nagar": "Rajendra Nagar",
    "Udyog Bhawan": "Seva Teerth",
    "Airport": "IGI Airport",
    "Yashoboomi Dwarka Sector - 25": "Yashobhoomi Dwarka Sector 25",
    "Phase 1": "DLF Phase 1",
    "Phase 2": "DLF Phase 2",
    "Phase 3": "Phase 3 Station",
    "Moulsari Avenue": "Moulsari Avenue Station",
    "Dwarka": "Dwarka - Kakrola",
    "Badkal Mor": "Badhkal Mor",
    "Sant Surdas Sihi": "Sant Surdas (Sihi)",
}
GATE_CODE_ALIASES = {"KG": "KGM", "AIIMS": "AIIM"}  # gate sheet code -> station sheet code

QUERY = """[out:json][timeout:150];
rel["route"~"subway|light_rail|monorail"]["network"~"Delhi Metro|Rapid Metro",i]->.r;
.r out body;
node(r.r)->.n;
.n out body;"""
MIRRORS = ["https://overpass-api.de/api/interpreter",
           "https://overpass.kumi.systems/api/interpreter"]


# --- helpers --------------------------------------------------------------------------

def fetch_osm():
    for url in MIRRORS:  # public Overpass servers 504 under load; try the next one
        req = urllib.request.Request(
            url,
            data=urllib.parse.urlencode({"data": QUERY}).encode(),
            headers={"User-Agent": "DPCS-course-project/0.1 (Delhi Metro route planner)"},
        )
        try:
            with urllib.request.urlopen(req, timeout=180) as r:
                OSM_RAW.write_bytes(r.read())
            return
        except OSError as e:
            print(f"{url}: {e}", file=sys.stderr)
    sys.exit("all Overpass mirrors failed; retry later")


def canonical(name):
    name = re.sub(r"\s*\((\w+ )?Line\)$", "", name.strip())  # "Welcome (Pink Line)" -> "Welcome"
    return ALIASES.get(name, name)


def norm(name):
    """Spelling-insensitive key: 'Noida Sec -18' == 'Noida Sector 18', drops '(Rapid Metro)'."""
    n = re.sub(r"\(.*?\)", "", str(name).lower())
    n = re.sub(r"\bsec\b", "sector", n)
    n = re.sub(r"\bmetro station\b", "", n)
    return re.sub(r"[^a-z0-9]", "", n)


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def km(a, b):
    p1, p2 = math.radians(a[0]), math.radians(b[0])
    dp, dl = p2 - p1, math.radians(b[1] - a[1])
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * 6371 * math.asin(math.sqrt(h))


def make_code(name, taken):
    """Fallback code for stations with no official code: initials, padded from the first word."""
    words = re.findall(r"[A-Za-z0-9]+", name)
    base = "".join(w[0] for w in words).upper()
    if len(base) < 3:
        base = (base + words[0][1:].upper())[:3]
    base, code, i = base[:4], base[:4], 2
    while code in taken:
        code, i = f"{base}{i}", i + 1
    return code


class Matcher:
    """Find the OSM station for an outside name/position: exact name, normalised name, nearby."""

    def __init__(self, stations):
        self.stations = stations
        self.by_norm = {norm(s["name"]): s for s in stations}

    def find(self, name, pt=None):
        s = self.by_norm.get(norm(name))
        if s:
            return s, "name"
        if pt:
            best = min(self.stations, key=lambda o: km(pt, (o["lat"], o["lon"])))
            if km(pt, (best["lat"], best["lon"])) <= MATCH_RADIUS_KM:
                return best, "position"
        return None, None


# --- 1. OSM: structure ----------------------------------------------------------------

def build_osm(raw):
    nodes = {e["id"]: e for e in raw["elements"] if e["type"] == "node"}
    rels = [e for e in raw["elements"] if e["type"] == "relation"]

    lines, seqs = {}, []
    pts = collections.defaultdict(list)          # station name -> [(lat, lon)]
    refs = collections.defaultdict(collections.Counter)
    st_lines = collections.defaultdict(set)
    for r in sorted(rels, key=lambda r: r["id"]):
        prefix = next((p for p in LINES if r["tags"].get("name", "").startswith(p)), None)
        if prefix is None:
            print("skipping unknown route:", r["tags"].get("name"), file=sys.stderr)
            continue
        lid, lname = LINES[prefix]
        lid, lname = next((v for (base, part), v in BRANCHES.items()
                           if base == lid and part in r["tags"]["name"]), (lid, lname))
        colour = r["tags"].get("colour", "#888888")
        lines[lid] = {"id": lid, "name": lname, "color": COLOR_FIX.get(colour.lower(), colour.upper())}
        seq = []
        for m in r["members"]:
            if m["type"] != "node" or not m["role"].startswith("stop") or m["ref"] not in nodes:
                continue
            n = nodes[m["ref"]]
            name = canonical(n["tags"]["name"])
            pts[name].append((n["lat"], n["lon"]))
            if n["tags"].get("ref"):
                refs[name][n["tags"]["ref"].upper()] += 1
            st_lines[name].add(lid)
            if not seq or seq[-1] != name:
                seq.append(name)
        seqs.append((lid, seq))

    stations = []
    for name, p in sorted(pts.items()):
        stations.append({"id": slug(name), "name": name,
                         "lat": round(sum(x[0] for x in p) / len(p), 6),
                         "lon": round(sum(x[1] for x in p) / len(p), 6),
                         "lines": sorted(st_lines[name]),
                         "_osm_ref": refs[name].most_common(1)[0][0] if refs[name] else None})
    edges = {}
    for lid, seq in seqs:
        for a, b in itertools.pairwise(seq):
            key = (lid, *sorted((slug(a), slug(b))))
            edges[key] = {"a": key[1], "b": key[2], "line": lid}
    return {"lines": sorted(lines.values(), key=lambda l: l["id"]),
            "stations": stations,
            "edges": sorted(edges.values(), key=lambda e: (e["line"], e["a"], e["b"]))}


# --- 2. GTFS: official ride times -----------------------------------------------------

def gtfs_times(matcher, report):
    """{frozenset(osm_id_a, osm_id_b): minutes}: median ride time (departure -> next arrival).

    Dwell is NOT added: ride times alone reproduce DMRC's planner (Rajiv Chowk -> Kashmere
    Gate: 8.7 min vs DMRC 8:37); adding the 20 s dwell per stop overshot by 1.4 min.
    """
    def rows(f):
        with open(GTFS / f, encoding="utf-8-sig") as fh:
            return list(csv.DictReader(fh))

    def sec(t):
        h, m, s = map(int, t.split(":"))
        return h * 3600 + m * 60 + s

    excluded = {r["route_id"] for r in rows("routes.txt")
                if r["route_long_name"].upper().startswith(GTFS_EXCLUDE_PREFIX)}
    trips = {t["trip_id"] for t in rows("trips.txt") if t["route_id"] not in excluded}
    used_stops = set()
    by_trip = collections.defaultdict(list)
    for r in rows("stop_times.txt"):
        if r["trip_id"] in trips:
            by_trip[r["trip_id"]].append(r)
            used_stops.add(r["stop_id"])

    to_osm, how = {}, collections.Counter()
    for s in rows("stops.txt"):
        if s["stop_id"] not in used_stops:
            continue
        found, via = matcher.find(s["stop_name"], (float(s["stop_lat"]), float(s["stop_lon"])))
        how[via or "unmatched"] += 1
        if found:
            to_osm[s["stop_id"]] = found["id"]
            if via == "position":
                report.append(f"  gtfs≈osm by position: {s['stop_name']!r} -> {found['name']!r}")
        else:
            report.append(f"  gtfs stop unmatched: {s['stop_name']!r}")
    report.append(f"GTFS stops: {dict(how)}")

    ride, dwell = collections.defaultdict(list), []
    for stop_rows in by_trip.values():
        stop_rows.sort(key=lambda r: int(r["stop_sequence"]))
        for p, n in itertools.pairwise(stop_rows):
            if p["stop_id"] in to_osm and n["stop_id"] in to_osm:
                ride[frozenset((to_osm[p["stop_id"]], to_osm[n["stop_id"]]))].append(
                    sec(n["arrival_time"]) - sec(p["departure_time"]))
        dwell += [sec(r["departure_time"]) - sec(r["arrival_time"]) for r in stop_rows]
    dwell_min = statistics.median(dwell) / 60
    report.append(f"GTFS median dwell {dwell_min * 60:.0f} s (not added, see docstring)")
    return {k: round(statistics.median(v) / 60, 2) for k, v in ride.items()}


# --- 3. Spreadsheet: codes and gates --------------------------------------------------

def parse_coord(v, hemi):
    m = re.fullmatch(rf"\s*(\d+\.\d+)\s*{hemi}?\s*", str(v))
    return float(m.group(1)) if m else None


CODE_RE = re.compile(r"[A-Z][A-Z0-9]{1,5}(?:[-\s]?\d+)?")  # RCK, KJMD-6, DW9; not "Phase 1"


def base_code(code):
    """Leading letters: 'CTST-2' -> 'CTST'. Only a lookup hint: 'JLA-8' is not 'JLA'."""
    return re.match(r"[A-Z]*", code).group()


def spreadsheet(matcher, report):
    """({osm_id: official code}, {osm_id: [gate dicts]})"""
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    st_rows = [r for r in list(wb["stations"].iter_rows(values_only=True))[1:] if any(r)]
    gate_rows = [r for r in list(wb["gates"].iter_rows(values_only=True))[1:] if any(r)]

    codes, code_to_id = collections.defaultdict(set), {}
    base_ids = collections.defaultdict(set)
    for line, code, name, *_rest in st_rows:
        lat, lon = parse_coord(_rest[2], "N"), parse_coord(_rest[3], "E")
        name = XLSX_ALIASES.get(str(name).strip(), name)
        found, _ = matcher.find(name, (lat, lon) if lat and lon else None)
        if not found:
            report.append(f"  spreadsheet station unmatched: {name!r} ({code})")
            continue
        code = str(code).strip()
        if not CODE_RE.fullmatch(code):  # Rapid Metro rows carry names in the code column
            continue
        codes[found["id"]].add(code)
        code_to_id[code] = found["id"]
        base_ids[base_code(code)].add(found["id"])

    # one station may have a code per line (AVIT, AVIT-7): prefer the plain one
    official = {sid: min(cs, key=lambda c: (bool(re.search(r"[^A-Z]", c)), len(c), c))
                for sid, cs in codes.items()}

    gates, skipped = collections.defaultdict(list), collections.Counter()
    for code, gate, location, *_ in gate_rows:
        if not code or not gate:
            skipped["no station code"] += 1
            continue
        code = GATE_CODE_ALIASES.get(str(code).strip().upper(), str(code).strip().upper())
        # variants like "CTST-2", "INA 7", "BCGN L-8": fall back to the leading letters
        sid = code_to_id.get(code)
        if not sid and len(base_ids[base_code(code)]) == 1:
            sid = next(iter(base_ids[base_code(code)]))
        if not sid:
            skipped["unknown station code"] += 1
            continue
        g = {"gate": str(gate).strip(), "location": " ".join(str(location or "").split())}
        if g not in gates[sid]:
            gates[sid].append(g)
    report.append(f"spreadsheet: codes for {len(official)} stations; gates for {len(gates)} stations "
                  f"({sum(map(len, gates.values()))} gates); skipped {dict(skipped)}")
    return official, gates


# --- 4. merge -------------------------------------------------------------------------

def fit_minutes_per_km(data, timed, report):
    """Least squares through the origin, minutes = k * km, on a fit set; checked on a hold-out."""
    coord = {s["id"]: (s["lat"], s["lon"]) for s in data["stations"]}
    groups = collections.defaultdict(lambda: {"fit": [], "hold": []})
    for e in data["edges"]:
        key = frozenset((e["a"], e["b"]))
        if key in timed:
            # deterministic split: same edge always lands in the same set
            part = "hold" if zlib.crc32(f"{e['a']}|{e['b']}".encode()) % 100 < HOLDOUT_SHARE * 100 else "fit"
            grp = e["line"] if e["line"] in SEPARATE_SPEED_LINES else "ALL"
            groups[grp][part].append((km(coord[e["a"]], coord[e["b"]]), timed[key]))
    k = {}
    for grp, sets in groups.items():
        fit = sets["fit"] or sets["hold"]
        k[grp] = sum(d * t for d, t in fit) / sum(d * d for d, _ in fit)
        errs = [abs(k[grp] * d - t) / t for d, t in sets["hold"]]
        report.append(f"estimate fit {grp}: {60 / k[grp]:.1f} km/h straight-line equivalent; "
                      f"fit n={len(sets['fit'])}, hold-out n={len(errs)}, "
                      f"hold-out mean abs error {100 * statistics.mean(errs) if errs else float('nan'):.0f}%")
    return k


def merge(data, timed, official, gates, report):
    coord = {s["id"]: (s["lat"], s["lon"]) for s in data["stations"]}
    k = fit_minutes_per_km(data, timed, report)
    for e in data["edges"]:
        key = frozenset((e["a"], e["b"]))
        if key in timed:
            e["minutes"], e["time_source"] = max(MIN_MINUTES, timed[key]), "gtfs"
        else:
            kk = k[e["line"] if e["line"] in k else "ALL"]
            e["minutes"] = max(MIN_MINUTES, round(kk * km(coord[e["a"]], coord[e["b"]]), 2))
            e["time_source"] = "estimate"
            report.append(f"  estimated (no GTFS): {e['line']} {e['a']} – {e['b']}")

    taken = set(official.values())
    source = collections.Counter()
    for s in data["stations"]:
        osm_ref = s.pop("_osm_ref")
        if s["id"] in official:
            s["code"] = official[s["id"]]
            source["official"] += 1
        elif osm_ref and osm_ref not in taken:   # OSM refs can be wrong (OSM says Bhalaswa = BIMR)
            s["code"] = osm_ref
            source["osm"] += 1
        else:
            s["code"] = make_code(s["name"], taken)
            source["generated"] += 1
        taken.add(s["code"])
        if gates.get(s["id"]):
            s["exits"] = gates[s["id"]]
    report.append(f"station codes: {dict(source)}")
    return [s["code"] for s in data["stations"] if s["id"] not in official]


def report_network(data, report):
    adj = collections.defaultdict(set)
    for e in data["edges"]:
        adj[e["a"]].add(e["b"])
        adj[e["b"]].add(e["a"])
    seen, stack = set(), [data["stations"][0]["id"]]
    while stack:
        v = stack.pop()
        if v not in seen:
            seen.add(v)
            stack.extend(adj[v] - seen)
    src = collections.Counter(e["time_source"] for e in data["edges"])
    report.append(f"stations {len(data['stations'])} · lines {len(data['lines'])} · edges {len(data['edges'])} "
                  f"· interchanges {sum(len(s['lines']) > 1 for s in data['stations'])} "
                  f"· connected {len(seen) == len(data['stations'])} · edge times {dict(src)}")


def main():
    if "--refresh" in sys.argv or not OSM_RAW.exists():
        fetch_osm()
    if not GTFS.exists() or not XLSX.exists():
        sys.exit(f"missing raw data in {HERE / 'raw'} (not in git: download DMRC GTFS and the station "
                 "& gate locations file from delhi.transportstack.in, see PLAN.md D4)")
    data = build_osm(json.loads(OSM_RAW.read_text()))
    report = []
    matcher = Matcher(data["stations"])
    timed = gtfs_times(matcher, report)
    official, gates = spreadsheet(matcher, report)
    non_official = merge(data, timed, official, gates, report)
    data["meta"] = {
        "sources": [
            "Stations, lines, coordinates: © OpenStreetMap contributors, ODbL 1.0",
            ("Travel times: DMRC static GTFS (Aug 2023); station codes and gates: DMRC station & gate "
            "locations. Both via Delhi Transport Stack / Delhi Open Transit Data, Transport Department, GNCTD"),
        ],
        "travel_time": "GTFS median ride time (no dwell) where available (time_source 'gtfs'); "
                       "otherwise straight-line distance x minutes-per-km fitted on GTFS times ('estimate')",
        "non_official_codes": sorted(non_official),
    }
    report_network(data, report)
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n")
    print("\n".join(report))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
